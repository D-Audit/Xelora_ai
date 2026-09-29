"""Bulk-load a user's stored CSV/TSV/XLSX file into the task workbook.

Moving a data file through the model (read_workspace_file -> write_table) is
slow, capped by tool payload size, and error-prone for more than a few
hundred rows. This tool reads the file on the backend and writes it straight
into Excel in chunks, then reads it back to verify.

Scope: only files inside the signed-in user's ``STORAGE_DIR/<user_id>``
(same resolver as the file tools). Values that would become formulas in Excel
(``=``, ``+``, ``-``, ``@`` prefixes that are not numbers) are written as text,
so an untrusted file cannot inject formulas such as WEBSERVICE or HYPERLINK.
"""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from typing import Any

MAX_IMPORT_ROWS = 200_000
MAX_IMPORT_COLUMNS = 300
_CHUNK_ROWS = 5_000
_NUMBER = re.compile(r"^[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?$")
_ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_ISO_DATETIME = re.compile(r"^\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}(:\d{2})?$")
_FORMULA_PREFIXES = ("=", "+", "-", "@")
SUPPORTED_EXTENSIONS = {".csv", ".tsv", ".txt", ".xlsx", ".xlsm"}


def coerce_cell(text: Any) -> Any:
    """Turn CSV text into the value Excel should store, never a formula."""
    if not isinstance(text, str):
        return text
    value = text.strip()
    if value == "":
        return None
    digits = sum(ch.isdigit() for ch in value)
    identifier_like = (
        (len(value) > 1 and value[0] == "0" and value[1].isdigit())   # ZIP codes, account numbers
        or value.startswith("+")                                       # phone numbers
        or (_NUMBER.match(value) is not None and digits > 15)          # beyond Excel's precision
    )
    if identifier_like:
        # Excel would reparse numeric-looking text; the apostrophe keeps it verbatim.
        return "'" + value
    if _NUMBER.match(value):
        try:
            number = float(value)
            return int(number) if number.is_integer() and "." not in value and "e" not in value.lower() else number
        except ValueError:
            pass
    if _ISO_DATE.match(value):
        try:
            return datetime.combine(date.fromisoformat(value), datetime.min.time())
        except ValueError:
            pass
    if _ISO_DATETIME.match(value):
        try:
            return datetime.fromisoformat(value.replace(" ", "T"))
        except ValueError:
            pass
    if value.startswith(_FORMULA_PREFIXES):
        return "'" + text  # stored as literal text by Excel
    return text


def parse_delimited(raw: bytes, suffix: str) -> list[list[Any]]:
    text = raw.decode("utf-8-sig", errors="replace")
    if suffix == ".tsv":
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(text[:20_000], delimiters=",;\t|").delimiter
        except csv.Error:
            # The sniffer gives up on ragged samples; the header line decides.
            header = text.splitlines()[0] if text else ""
            delimiter = max(",;\t|", key=header.count) if any(ch in header for ch in ",;\t|") else ","
    rows = []
    for index, row in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter)):
        if index >= MAX_IMPORT_ROWS + 1:
            raise ValueError(f"The file has more than {MAX_IMPORT_ROWS} rows; split it before importing.")
        rows.append([coerce_cell(cell) for cell in row[:MAX_IMPORT_COLUMNS]])
    return rows


def parse_workbook(path: str, sheet: str | None) -> list[list[Any]]:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet and sheet not in book.sheetnames:
            raise ValueError(f"Sheet '{sheet}' is not in the file. Sheets: {book.sheetnames}.")
        worksheet = book[sheet] if sheet else book.worksheets[0]
        rows = []
        for index, row in enumerate(worksheet.iter_rows(values_only=True)):
            if index >= MAX_IMPORT_ROWS + 1:
                raise ValueError(f"The sheet has more than {MAX_IMPORT_ROWS} rows; split it before importing.")
            rows.append([coerce_cell(value) if isinstance(value, str) else value for value in row[:MAX_IMPORT_COLUMNS]])
        return rows
    finally:
        book.close()


def rectangular(rows: list[list[Any]]) -> list[list[Any]]:
    """Pad ragged rows and drop trailing empty rows so Excel gets one block."""
    while rows and all(value in (None, "") for value in rows[-1]):
        rows.pop()
    width = max((len(row) for row in rows), default=0)
    return [row + [None] * (width - len(row)) for row in rows]


def load_rows(path: str, suffix: str, sheet: str | None) -> list[list[Any]]:
    if suffix in {".xlsx", ".xlsm"}:
        return rectangular(parse_workbook(path, sheet))
    with open(path, "rb") as handle:
        return rectangular(parse_delimited(handle.read(), suffix))


def import_workspace_file_to_sheet(
    user_id: int | None,
    path: str,
    sheet_name: str,
    start_cell: str = "A1",
    source_sheet: str | None = None,
    overwrite: bool = False,
    *,
    workbook_name: str | None = None,
    excel_app_pid: int | None = None,
) -> dict[str, Any]:
    from tools.filesystem import _safe_existing_path

    root, target = _safe_existing_path(user_id, path)
    suffix = target.suffix.lower()
    if not target.is_file() or suffix not in SUPPORTED_EXTENSIONS:
        return {"verified": False, "status": "unsupported_file",
                "error": f"Import supports {', '.join(sorted(SUPPORTED_EXTENSIONS))} files in your Xelora storage."}
    try:
        rows = load_rows(str(target), suffix, source_sheet)
    except (ValueError, OSError) as exc:
        return {"verified": False, "status": "unreadable_file", "error": str(exc)}
    except Exception as exc:  # openpyxl raises several types for corrupt files
        return {"verified": False, "status": "unreadable_file", "error": f"Could not read the workbook: {type(exc).__name__}: {exc}"}
    if not rows:
        return {"verified": False, "status": "empty_file", "error": "The file contains no data."}

    from skills.excel_shared import bind_workbook_context, get_active_workbook
    if workbook_name:
        bind_workbook_context(workbook_name, excel_app_pid)
    wb = get_active_workbook()
    if sheet_name not in [sheet.name for sheet in wb.sheets]:
        wb.sheets.add(sheet_name, after=wb.sheets[-1])
    sheet = wb.sheets[sheet_name]
    height, width = len(rows), len(rows[0])
    destination = sheet.range(start_cell).resize(height, width)
    if not overwrite:
        existing = destination.options(ndim=2).value or []
        if any(value not in (None, "") for row in existing for value in row):
            return {"verified": False, "status": "target_not_empty",
                    "error": f"'{sheet_name}'!{destination.address} already has data. Choose an empty sheet/cell or pass overwrite=true."}
    anchor = sheet.range(start_cell)
    for offset in range(0, height, _CHUNK_ROWS):
        chunk = rows[offset:offset + _CHUNK_ROWS]
        anchor.offset(offset, 0).resize(len(chunk), width).value = chunk
    wb.save()

    first = destination.resize(1, width).options(ndim=2).value[0]
    last = destination.offset(height - 1, 0).resize(1, width).options(ndim=2).value[0]

    def same(expected: list[Any], actual: list[Any]) -> bool:
        for want, got in zip(expected, actual):
            if isinstance(want, str) and want.startswith("'"):
                want = want[1:]
            if want in (None, "") and got in (None, ""):
                continue
            if isinstance(want, (int, float)) and isinstance(got, (int, float)):
                if abs(float(want) - float(got)) > 1e-9 * max(1.0, abs(float(want))):
                    return False
                continue
            if isinstance(want, datetime) and isinstance(got, datetime):
                if want.replace(microsecond=0) != got.replace(microsecond=0):
                    return False
                continue
            if str(want) != str(got):
                return False
        return True

    verified = same(rows[0], first) and same(rows[-1], last)
    return {
        "source": str(target.relative_to(root)), "sheet": sheet_name,
        "range": f"'{sheet_name}'!{destination.address}", "rows_written": height, "columns": width,
        "header_row": [str(value) for value in rows[0]][:50],
        "status": "imported" if verified else "import_mismatch", "verified": verified,
        "verification_note": (
            f"Read back the first and last of {height} rows and matched the file."
            if verified else "First/last rows read back from Excel do not match the file."
        ),
        "next_steps": "Use write_table with rows=[] and a table_name to make it a native Excel Table if needed.",
    }
